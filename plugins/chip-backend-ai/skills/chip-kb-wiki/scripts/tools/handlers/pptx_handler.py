"""pptx handler: scan slide structure and parse a slide range."""

from __future__ import annotations

from dataclasses import dataclass, field
from pathlib import Path

from pptx import Presentation
from pptx.enum.shapes import MSO_SHAPE_TYPE


@dataclass
class SlideMeta:
    idx: int
    title: str
    images: int
    has_notes: bool


@dataclass
class PptxScan:
    slides: list[SlideMeta] = field(default_factory=list)


def _slide_title(slide) -> str:
    if slide.shapes.title is not None:
        text = (slide.shapes.title.text or "").strip()
        if text:
            return text
    for shape in slide.shapes:
        if shape.has_text_frame:
            text = (shape.text_frame.text or "").strip()
            if text:
                return text.split("\n", 1)[0][:80]
    return ""


def _count_images(slide) -> int:
    count = 0
    for shape in slide.shapes:
        if shape.shape_type == MSO_SHAPE_TYPE.PICTURE:
            count += 1
        elif shape.shape_type == MSO_SHAPE_TYPE.GROUP:
            count += _count_group_images(shape)
    return count


def _count_group_images(group) -> int:
    count = 0
    for shape in group.shapes:
        if shape.shape_type == MSO_SHAPE_TYPE.PICTURE:
            count += 1
        elif shape.shape_type == MSO_SHAPE_TYPE.GROUP:
            count += _count_group_images(shape)
    return count


def _has_notes(slide) -> bool:
    if not slide.has_notes_slide:
        return False
    notes = slide.notes_slide.notes_text_frame.text if slide.notes_slide else ""
    return bool(notes and notes.strip())


def scan_meta(path: Path) -> PptxScan:
    prs = Presentation(str(path))
    slides: list[SlideMeta] = []
    for i, slide in enumerate(prs.slides, start=1):
        slides.append(
            SlideMeta(
                idx=i,
                title=_slide_title(slide) or f"Slide {i}",
                images=_count_images(slide),
                has_notes=_has_notes(slide),
            )
        )
    return PptxScan(slides=slides)


def _shape_text_lines(shape) -> list[str]:
    if not shape.has_text_frame:
        return []
    lines: list[str] = []
    for para in shape.text_frame.paragraphs:
        text = "".join(run.text for run in para.runs).strip()
        if text:
            lines.append(text)
    return lines


def _shape_table_md(shape) -> str:
    table = shape.table
    rows = [
        [cell.text.strip().replace("\n", " ") for cell in row.cells]
        for row in table.rows
    ]
    if not rows:
        return ""
    cols = max(len(r) for r in rows)
    rows = [r + [""] * (cols - len(r)) for r in rows]
    header = rows[0]
    body = rows[1:] if len(rows) > 1 else []
    lines = ["| " + " | ".join(header) + " |", "| " + " | ".join(["---"] * cols) + " |"]
    for r in body:
        lines.append("| " + " | ".join(r) + " |")
    return "\n".join(lines)


def _slide_to_markdown(slide, idx: int) -> str:
    title = _slide_title(slide) or f"Slide {idx}"
    out = [f"## Slide {idx} · {title}", ""]

    img_counter = 0
    for shape in slide.shapes:
        if shape.shape_type == MSO_SHAPE_TYPE.PICTURE or (
            shape.shape_type == MSO_SHAPE_TYPE.GROUP and _count_group_images(shape) > 0
        ):
            n = (
                1
                if shape.shape_type == MSO_SHAPE_TYPE.PICTURE
                else _count_group_images(shape)
            )
            for _ in range(n):
                img_counter += 1
                alt = (shape.name or "image").strip()
                out.append(f"> [图占位 · slide={idx} · img#{img_counter} · name={alt}]")
            continue
        if shape.has_table:
            md = _shape_table_md(shape)
            if md:
                out.append("")
                out.append(md)
                out.append("")
            continue
        for line in _shape_text_lines(shape):
            if line == title and out[0].endswith(title):
                continue
            out.append(line)

    if slide.has_notes_slide:
        notes = slide.notes_slide.notes_text_frame.text.strip() if slide.notes_slide else ""
        if notes:
            out.append("")
            out.append("**演讲者备注：**")
            out.append("")
            out.append(notes)

    return "\n".join(out).rstrip() + "\n"


def parse_slides(path: Path, start: int, end: int) -> str:
    prs = Presentation(str(path))
    total = len(prs.slides)
    if start < 1 or end < start or start > total:
        raise ValueError(f"invalid slide range {start}-{end} (have 1..{total})")
    end = min(end, total)
    out: list[str] = [f"# Slides {start}..{end} of {path.name}", ""]
    for i, slide in enumerate(prs.slides, start=1):
        if i < start:
            continue
        if i > end:
            break
        out.append(_slide_to_markdown(slide, i))
    return "\n".join(out)
