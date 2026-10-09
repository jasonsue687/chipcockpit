"""xlsx handler: scan sheet structure and parse a sheet (full / schema / row range)."""

from __future__ import annotations

from dataclasses import dataclass, field
from pathlib import Path

from openpyxl import load_workbook
from openpyxl.worksheet.worksheet import Worksheet


DENSE_THRESHOLD = 5000
SCHEMA_PREVIEW_ROWS = 5


@dataclass
class SheetMeta:
    name: str
    rows: int
    cols: int
    has_header: bool
    dense: bool


@dataclass
class XlsxScan:
    sheets: list[SheetMeta] = field(default_factory=list)


def _open_readonly(path: Path):
    return load_workbook(filename=str(path), read_only=True, data_only=True)


def _effective_dims(ws: Worksheet) -> tuple[int, int]:
    rows = ws.max_row or 0
    cols = ws.max_column or 0
    return rows, cols


def _looks_like_header(first_row: list) -> bool:
    if not first_row:
        return False
    str_cells = sum(1 for c in first_row if isinstance(c, str) and c.strip())
    return str_cells >= max(1, len(first_row) // 2)


def scan_meta(path: Path) -> XlsxScan:
    wb = _open_readonly(path)
    sheets: list[SheetMeta] = []
    try:
        for name in wb.sheetnames:
            ws = wb[name]
            rows, cols = _effective_dims(ws)
            first_row: list = []
            if rows > 0:
                row_iter = ws.iter_rows(min_row=1, max_row=1, values_only=True)
                first_row = list(next(row_iter, ())) or []
            sheets.append(
                SheetMeta(
                    name=name,
                    rows=rows,
                    cols=cols,
                    has_header=_looks_like_header(first_row),
                    dense=rows >= DENSE_THRESHOLD,
                )
            )
    finally:
        wb.close()
    return XlsxScan(sheets=sheets)


def _cell_to_str(value) -> str:
    if value is None:
        return ""
    if isinstance(value, str):
        return value.strip().replace("\n", " ").replace("|", "\\|")
    return str(value)


def _rows_to_markdown(rows: list[list], has_header: bool) -> str:
    if not rows:
        return "_(空)_"
    cols = max(len(r) for r in rows)
    rows = [list(r) + [""] * (cols - len(r)) for r in rows]
    str_rows = [[_cell_to_str(c) for c in r] for r in rows]
    if has_header and len(str_rows) >= 1:
        header = str_rows[0]
        body = str_rows[1:]
    else:
        header = [f"col{i + 1}" for i in range(cols)]
        body = str_rows
    lines = ["| " + " | ".join(header) + " |", "| " + " | ".join(["---"] * cols) + " |"]
    for r in body:
        lines.append("| " + " | ".join(r) + " |")
    return "\n".join(lines)


def _sheet_meta(path: Path, sheet_name: str) -> SheetMeta:
    for s in scan_meta(path).sheets:
        if s.name == sheet_name:
            return s
    raise ValueError(f"sheet not found: {sheet_name!r}")


def parse_sheet(
    path: Path,
    sheet_name: str,
    schema_only: bool = False,
    rows_range: tuple[int, int] | None = None,
) -> str:
    meta = _sheet_meta(path, sheet_name)
    wb = _open_readonly(path)
    try:
        if sheet_name not in wb.sheetnames:
            raise ValueError(f"sheet not found: {sheet_name!r}")
        ws = wb[sheet_name]
        out: list[str] = [
            f"# Sheet: {sheet_name}",
            "",
            f"- rows: {meta.rows} | cols: {meta.cols} | has_header: {meta.has_header} | dense: {meta.dense}",
            "",
        ]
        if meta.rows == 0:
            out.append("_(空 sheet)_")
            return "\n".join(out) + "\n"

        if schema_only:
            preview_n = min(SCHEMA_PREVIEW_ROWS + (1 if meta.has_header else 0), meta.rows)
            preview_rows = [
                list(r)
                for r in ws.iter_rows(min_row=1, max_row=preview_n, values_only=True)
            ]
            out.append("## Schema preview")
            out.append("")
            out.append(_rows_to_markdown(preview_rows, meta.has_header))
            out.append("")
            out.append(f"_(schema_only: 仅展示前 {preview_n} 行)_")
            return "\n".join(out) + "\n"

        if rows_range is not None:
            r_start, r_end = rows_range
            r_start = max(1, r_start)
            r_end = min(meta.rows, r_end)
            if r_start > r_end:
                raise ValueError(f"invalid rows range {rows_range}")
            data_rows = [
                list(r)
                for r in ws.iter_rows(
                    min_row=r_start, max_row=r_end, values_only=True
                )
            ]
            header_row: list = []
            if meta.has_header and r_start > 1:
                header_row = list(
                    next(ws.iter_rows(min_row=1, max_row=1, values_only=True), ())
                )
                rows_for_md = [header_row, *data_rows]
                out.append(f"## Rows {r_start}..{r_end}（含表头）")
            else:
                rows_for_md = data_rows
                out.append(f"## Rows {r_start}..{r_end}")
            out.append("")
            out.append(_rows_to_markdown(rows_for_md, meta.has_header))
            return "\n".join(out) + "\n"

        if meta.dense:
            return _build_schema_only_with_warning(out, ws, meta)

        all_rows = [list(r) for r in ws.iter_rows(values_only=True)]
        out.append("## Full sheet")
        out.append("")
        out.append(_rows_to_markdown(all_rows, meta.has_header))
        return "\n".join(out) + "\n"
    finally:
        wb.close()


def _build_schema_only_with_warning(
    out: list[str], ws: Worksheet, meta: SheetMeta
) -> str:
    preview_n = min(SCHEMA_PREVIEW_ROWS + (1 if meta.has_header else 0), meta.rows)
    preview_rows = [
        list(r)
        for r in ws.iter_rows(min_row=1, max_row=preview_n, values_only=True)
    ]
    out.append("## Schema preview (auto: dense sheet)")
    out.append("")
    out.append(
        f"> 该 sheet 共 {meta.rows} 行 ≥ {DENSE_THRESHOLD} 行阈值，仅返回前 {preview_n} 行。"
        " 如需具体数据请重新调用并指定 --rows START:END。"
    )
    out.append("")
    out.append(_rows_to_markdown(preview_rows, meta.has_header))
    return "\n".join(out) + "\n"
