"""docx handler: scan structure (adaptive section split) and parse a single section.

Adaptive split rule:
1. Find the smallest heading level present that has >=2 occurrences -> split level L
2. If every present heading level has only 1 occurrence -> use the smallest present
3. Any heading at level <= L starts a new section (so a single H1 above many H2s
   becomes its own section, then H2s split the rest)
4. If no headings at all -> the whole doc is one section titled by filename
"""

from __future__ import annotations

from dataclasses import dataclass, field
from pathlib import Path
from typing import Iterator

from docx import Document
from docx.document import Document as DocxDocument
from docx.oxml.ns import qn
from docx.table import Table
from docx.text.paragraph import Paragraph


@dataclass
class SectionMeta:
    idx: int
    title: str
    paragraphs: int = 0
    images: int = 0
    tables: int = 0


@dataclass
class DocxScan:
    sections: list[SectionMeta] = field(default_factory=list)
    split_level: int | None = None  # the heading level used to split; None = no headings


def _iter_blocks(doc: DocxDocument) -> Iterator[Paragraph | Table]:
    """Yield paragraphs and tables in document order."""
    body = doc.element.body
    for child in body.iterchildren():
        tag = child.tag
        if tag == qn("w:p"):
            yield Paragraph(child, doc)
        elif tag == qn("w:tbl"):
            yield Table(child, doc)


def _heading_level(p: Paragraph) -> int | None:
    style = (p.style.name or "") if p.style else ""
    if not style.startswith("Heading"):
        return None
    rest = style[len("Heading") :].strip()
    try:
        return int(rest)
    except ValueError:
        return None


def _count_drawings(p: Paragraph) -> int:
    return len(p._p.findall(f".//{qn('w:drawing')}"))


def _detect_split_level(blocks: list) -> int | None:
    counts: dict[int, int] = {}
    for b in blocks:
        if isinstance(b, Paragraph):
            lvl = _heading_level(b)
            if lvl is not None:
                counts[lvl] = counts.get(lvl, 0) + 1
    if not counts:
        return None
    multi = [lvl for lvl, n in counts.items() if n >= 2]
    if multi:
        return min(multi)
    return min(counts.keys())


def _is_section_break(p: Paragraph, split_level: int) -> bool:
    lvl = _heading_level(p)
    return lvl is not None and lvl <= split_level


def scan_meta(path: Path) -> DocxScan:
    doc = Document(str(path))
    blocks = list(_iter_blocks(doc))
    split_level = _detect_split_level(blocks)
    sections: list[SectionMeta] = []

    if split_level is None:
        sec = SectionMeta(idx=1, title=path.stem)
        for block in blocks:
            _accumulate(block, sec)
        sections.append(sec)
        return DocxScan(sections=sections, split_level=None)

    current: SectionMeta | None = None
    for block in blocks:
        if isinstance(block, Paragraph) and _is_section_break(block, split_level):
            current = SectionMeta(
                idx=len(sections) + 1,
                title=block.text.strip() or f"Section {len(sections) + 1}",
            )
            sections.append(current)
            continue
        if current is None:
            continue  # pre-first-heading content (e.g. cover page) is dropped from outline
        _accumulate(block, current)

    if not sections:
        sec = SectionMeta(idx=1, title=path.stem)
        for block in blocks:
            _accumulate(block, sec)
        sections.append(sec)

    return DocxScan(sections=sections, split_level=split_level)


def _accumulate(block, sec: SectionMeta) -> None:
    if isinstance(block, Paragraph):
        if block.text.strip():
            sec.paragraphs += 1
        sec.images += _count_drawings(block)
    else:
        sec.tables += 1
        for row in block.rows:
            for cell in row.cells:
                for p in cell.paragraphs:
                    sec.images += _count_drawings(p)


def _table_to_markdown(table: Table) -> str:
    rows = [[cell.text.strip().replace("\n", " ") for cell in row.cells] for row in table.rows]
    if not rows:
        return ""
    cols = max(len(r) for r in rows)
    rows = [r + [""] * (cols - len(r)) for r in rows]
    header = rows[0]
    body = rows[1:] if len(rows) > 1 else []
    lines = ["| " + " | ".join(header) + " |", "| " + " | ".join(["---"] * cols) + " |"]
    for r in body:
        lines.append("| " + " | ".join(r) + " |")
    return "\n".join(lines)


def _paragraph_to_markdown(p: Paragraph, section_idx: int, para_idx: int) -> str:
    level = _heading_level(p)
    text = p.text.strip()
    drawings = _count_drawings(p)

    chunks: list[str] = []
    if level is not None and level >= 2:
        chunks.append("#" * (level + 1) + " " + (text or f"(heading {level})"))
    elif text:
        style_name = (p.style.name or "") if p.style else ""
        if style_name.startswith("List"):
            chunks.append(f"- {text}")
        else:
            chunks.append(text)
    for i in range(drawings):
        chunks.append(f"> [图占位 · section={section_idx} · paragraph={para_idx} · img#{i + 1}]")
    return "\n".join(chunks)


def parse_section(path: Path, section_idx: int) -> str:
    doc = Document(str(path))
    blocks = list(_iter_blocks(doc))
    scan = scan_meta(path)
    if section_idx < 1 or section_idx > len(scan.sections):
        raise ValueError(
            f"section {section_idx} out of range (have 1..{len(scan.sections)})"
        )

    target = scan.sections[section_idx - 1]
    out: list[str] = [f"# {target.title}", ""]
    para_idx = 0

    if scan.split_level is None:
        for block in blocks:
            para_idx = _emit_block(block, section_idx, para_idx, out)
        return "\n".join(out).rstrip() + "\n"

    cur_idx = 0
    inside = False
    for block in blocks:
        if isinstance(block, Paragraph) and _is_section_break(block, scan.split_level):
            cur_idx += 1
            if cur_idx > section_idx:
                break
            inside = cur_idx == section_idx
            continue
        if not inside:
            continue
        para_idx = _emit_block(block, section_idx, para_idx, out)

    return "\n".join(out).rstrip() + "\n"


def _emit_block(block, section_idx: int, para_idx: int, out: list[str]) -> int:
    if isinstance(block, Paragraph):
        para_idx += 1
        md = _paragraph_to_markdown(block, section_idx, para_idx)
        if md:
            out.append(md)
    else:
        md = _table_to_markdown(block)
        if md:
            out.append("")
            out.append(md)
            out.append("")
    return para_idx
